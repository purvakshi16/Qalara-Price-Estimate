from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = "Qalara Price Estimate Project Brief.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Aptos"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = header_cells[i]
        cell.text = header
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, "263A34")
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = value
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                shade_cell(cell, "F4F7F5")
            set_cell_border(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def build():
    doc = Document()
    style_doc(doc)

    doc.add_paragraph("Qalara Price Estimate Project Brief", style="Title")

    add_heading(doc, "Overview")
    doc.add_paragraph(
        "The Qalara Price Estimate application helps a user estimate the current market selling price of a product by combining product details, a reference image, and regional market context. The app searches for similar publicly listed products, extracts retail price evidence from real product detail pages, and presents a clean estimate with source rows that include marketplace, seller, product image, retail price, product link, and evidence-backed MOQ or wholesale details where available. The goal is to reduce manual price research while keeping the result traceable to source links."
    )

    add_heading(doc, "Features")
    add_table(
        doc,
        ["Feature", "What It Does"],
        [
            ["Product input form", "Collects product name, description, material, reference image, and target region."],
            ["Region and currency", "Defaults to US and USD, with options for UK GBP and EU EUR with Germany as the EU search basis."],
            ["Image upload", "Lets the user add a product reference image for visual context. Source rows use matched product images, not the uploaded image."],
            ["Live source search", "Uses the backend search agent to look for similar product listings and direct product detail pages."],
            ["Estimate cards", "Shows price range, single estimate, and confidence based on priced matches."],
            ["Source table", "Shows source product image, marketplace, seller, retail price, MOQ evidence, and product link."],
            ["MOQ evidence cell", "Shows suggested MOQ, wholesale price, and seller wholesale source only when product-relevant evidence is found."],
            ["Download menu", "Provides PDF and Excel CSV export options from a single Download button."],
        ],
        widths=[1.9, 4.55],
    )

    add_heading(doc, "Logic Behind The Features")
    add_table(
        doc,
        ["Area", "Logic"],
        [
            ["Retail price estimate", "Retail prices come from similar public product listings. The app uses multiple priced matches rather than relying on one source."],
            ["Single estimate", "The single estimate uses the median price so unusually high or low listings do not dominate the result."],
            ["Price range", "The range shows the lowest and highest retail prices from the valid priced matches used for the estimate."],
            ["Marketplace", "Marketplace means the platform or website hosting the product page, such as Amazon, eBay, Alibaba, Etsy, or a retailer site."],
            ["Seller", "Seller means the actual seller, store, brand, or vendor shown on the product detail page. It should not simply repeat the marketplace name."],
            ["Product link", "The product link must point to the exact product detail page used as retail price evidence, not a general search result page."],
            ["MOQ and wholesale price", "MOQ is only shown when the app finds product-specific or product-category-specific wholesale evidence from the seller, supplier, trade, bulk, or B2B page."],
            ["Not found behavior", "If seller wholesale evidence is not found, the MOQ cell shows Not found rather than guessing."],
            ["Amazon handling", "Amazon can appear only when it naturally appears in search results. The app does not use Amazon login credentials."],
            ["Accuracy principle", "The app favors evidence-backed fields over guessed values. Returning Not found is preferred when wholesale information cannot be verified."],
        ],
        widths=[1.85, 4.6],
    )

    add_heading(doc, "Current Search Rules")
    add_bullets(
        doc,
        [
            "Visit up to 15 candidate product links during live search.",
            "Keep up to 10 valid priced products for estimate calculation.",
            "Display the top 6 source rows in the estimate table.",
            "Reject general search pages as final retail sources.",
            "Keep retail source links and MOQ source links separate.",
        ],
    )

    add_heading(doc, "Expected Output")
    doc.add_paragraph(
        "For each source row, the user should be able to see what product was matched, where it was found, who the seller is, what retail price was listed, whether product-relevant wholesale information was found, and which source link supports each value."
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
