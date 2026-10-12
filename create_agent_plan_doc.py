from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = "Qalara Price Estimate Agent Plan.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        shade_cell(hdr[i], "263A34")
        for paragraph in hdr[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
        set_cell_border(hdr[i])
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            hdr[i].width = Inches(widths[i])

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if row_index % 2 == 1:
                shade_cell(cells[i], "F4F7F5")
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Inches(widths[i])
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()
    return table


def add_screenshot_placeholder(doc, title):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = f"Screenshot placeholder: {title}\nPaste screenshot here."
    shade_cell(cell, "F7F8F4")
    set_cell_border(cell, "BFC8C2")
    set_cell_margins(cell, top=220, bottom=220, start=220, end=220)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(90, 103, 98)
    doc.add_paragraph()


def build_doc():
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Qalara Price Estimate Agent Plan")

    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run(
        "This document explains what the price estimate application should do and how the agent should process a product search. "
        "It is written as a simple working plan so screenshots and implementation notes can be added later."
    )

    add_heading(doc, "Inputs")
    add_table(
        doc,
        ["Input", "Purpose", "Current Rule"],
        [
            ["Product name", "Primary search term and row label.", "Required."],
            ["Description", "Adds shape, finish, usage, pack size, or other product hints.", "Optional but useful."],
            ["Material", "Improves match quality and helps avoid unrelated products.", "Optional."],
            ["Region", "Controls search market and assumed currency.", "UK means GBP, US means USD, EU defaults to Germany and EUR."],
            ["Upload image", "Reference for image-led search and visual similarity.", "One image for now."],
        ],
        widths=[1.35, 3.0, 2.1],
    )

    add_heading(doc, "Agent Workflow")
    add_numbered(
        doc,
        [
            "Receive product name, description, material, region, and image reference.",
            "Generate a search task using the selected region.",
            "Search for visually and textually similar products from public web results.",
            "Visit up to 15 candidate product links.",
            "Reject general search pages, category pages, and pages without usable price data.",
            "Extract up to 10 valid priced products for the estimate calculation.",
            "Display the top 6 most useful source rows to the user.",
            "Let the user accept or reject source rows before relying on them.",
            "Calculate price range, median single estimate, and confidence.",
            "Allow PDF and Excel compatible CSV export.",
        ],
    )

    add_heading(doc, "Output Rows")
    doc.add_paragraph("Each source row should represent one real product listing and contain the fields below.")
    add_table(
        doc,
        ["Field", "Meaning"],
        [
            ["Source product image", "Image extracted from the matched product page."],
            ["Product title", "Title of the product on the source page."],
            ["Price and currency", "Visible selling price normalized to the region currency when possible."],
            ["Seller or store", "Seller name, store name, marketplace name, or best available source identity."],
            ["MOQ", "Minimum order quantity if visible. Otherwise show Not listed."],
            ["Exact product URL", "The direct product page link used as evidence."],
        ],
        widths=[1.8, 4.65],
    )

    add_heading(doc, "Estimate Logic")
    doc.add_paragraph(
        "The app should not estimate from only the 5 or 6 rows shown on screen. The backend should visit up to 15 candidate links, "
        "keep up to 10 valid priced product listings for calculation, and display the top 6 rows for review. "
        "The estimate should show a price range, median single estimate, and confidence label."
    )
    add_table(
        doc,
        ["Metric", "How It Is Used"],
        [
            ["Price range", "Minimum to maximum price from accepted or valid priced products."],
            ["Single estimate", "Median price, because it is less affected by unusually high or low listings."],
            ["Confidence", "Higher when more valid, similar, priced listings are found."],
        ],
        widths=[1.8, 4.65],
    )

    add_heading(doc, "Screenshot Areas")
    doc.add_paragraph("Use the spaces below to paste screenshots of the working app, search results, and any errors.")
    add_screenshot_placeholder(doc, "Input form with product image upload")
    add_screenshot_placeholder(doc, "Estimate panel and source rows")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
