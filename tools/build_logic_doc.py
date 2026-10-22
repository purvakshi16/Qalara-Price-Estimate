from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SEARCH-LOGIC.md"
OUTPUT = ROOT / "Qalara Price Estimate Agent Plan - Updated.docx"


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def clean(text):
    return re.sub(r"[`*_]", "", text).strip()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor(35, 39, 36)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for name, size, before, after in (("Heading 1", 16, 15, 6), ("Heading 2", 12, 10, 4)):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

if "Small" not in styles:
    small = styles.add_style("Small", WD_STYLE_TYPE.PARAGRAPH)
else:
    small = styles["Small"]
small.font.name = "Aptos"
small.font.size = Pt(9)
small.font.color.rgb = RGBColor(95, 98, 92)

header = section.header.paragraphs[0]
header.text = "Qalara Market Price Estimator"
header.style = styles["Small"]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Production search logic  |  ")
footer_run.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

title_style = styles.add_style("Document Title", WD_STYLE_TYPE.PARAGRAPH)
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(10)
title_style.paragraph_format.keep_with_next = True
title = doc.add_paragraph(style="Document Title")
title.add_run("Qalara Market Price Estimator Search Logic")
subtitle = doc.add_paragraph("Production rulebook and implementation record", style="Subtitle")
subtitle.runs[0].font.color.rgb = RGBColor(86, 77, 62)
doc.add_paragraph("Updated 22 October 2026", style="Small")
doc.add_paragraph(
    "This document is the detailed source of truth for how the application accepts product evidence, discovers comparable retail products, validates product pages, extracts prices and assembles results. It records both major workflow decisions and small thresholds so future changes can be reviewed without relying on memory."
)
doc.add_paragraph(
    "The current system estimates regional retail selling prices from live product pages. It does not perform wholesale research, store search history, train a model or claim a numerical visual-similarity score."
)

doc.add_heading("Current Production Scope", level=1)
scope = doc.add_table(rows=1, cols=2)
scope.style = "Table Grid"
scope.autofit = False
scope.columns[0].width = Inches(2.05)
scope.columns[1].width = Inches(4.65)
for index, value in enumerate(("Area", "Current behavior")):
    cell = scope.rows[0].cells[index]
    cell.text = value
    cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
    cell._tc.tcPr[-1].set(qn("w:fill"), "5F6F52")
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
rows = [
    ("Inputs", "Product name or image is required; either may be used alone or together. Description, material and dimensions are optional. Region is required."),
    ("Outputs", "Up to six direct product results with source image, linked product title, PDP price when verified, marketplace and seller."),
    ("Summary", "Price Range uses verified priced PDPs only. Diagnostics show visual matches returned, pages accessed and valid priced PDPs found."),
    ("Providers", "Selected SearchAPI or SerpApi for visual discovery; TinyFish for text discovery and page reading; limited Firecrawl discovery and fetch recovery when configured."),
    ("Storage", "No result database. Temporary public image upload is deleted after each request."),
]
for row_index, values in enumerate(rows, start=1):
    cells = scope.add_row().cells
    for column, value in enumerate(values):
        cells[column].text = value
        cells[column].vertical_alignment = 1
        if row_index % 2 == 0:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F4F2EC")
            cells[column]._tc.get_or_add_tcPr().append(shading)
for row in scope.rows:
    row.cells[0].width = Inches(2.05)
    row.cells[1].width = Inches(4.65)
    for cell in row.cells:
        set_cell_border(cell)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)

doc.add_heading("How To Maintain This Document", level=1)
doc.add_paragraph("Update this Word file and SEARCH-LOGIC.md in the same change whenever a provider, threshold, field, timeout, result rule or user-facing output changes. PRODUCTION-FEATURES.md remains the shorter feature summary.")

lines = SOURCE.read_text(encoding="utf-8").splitlines()
for line in lines:
    stripped = line.strip()
    if not stripped or stripped.startswith("# Market Price") or stripped.startswith("Last updated:") or stripped.startswith("This is the canonical"):
        continue
    if stripped.startswith("## "):
        doc.add_heading(clean(stripped[3:]), level=1)
    elif re.match(r"^\d+\. ", stripped):
        paragraph = doc.add_paragraph(style="Normal")
        number, value = stripped.split(". ", 1)
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.add_run(f"{number}.  {clean(value)}")
    elif stripped.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(clean(stripped[2:]))
    else:
        doc.add_paragraph(clean(stripped))

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True
    if paragraph.text == "13 Marketplace and Seller Logic":
        paragraph.paragraph_format.page_break_before = True
    if paragraph.style.name == "List Bullet":
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)

doc.core_properties.title = "Qalara Market Price Estimator Search Logic"
doc.core_properties.subject = "Production rules for product discovery, PDP validation and retail price extraction"
doc.core_properties.author = "Qalara"
doc.save(OUTPUT)
print(OUTPUT)
